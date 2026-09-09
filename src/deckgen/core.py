"""
deckgen — one slide spec, three outputs.

A deck is a plain Python structure (see the course repo's deck/*.py). Every slide is
laid out once, on a 1920 x 1080 px canvas, into a list of primitive
elements (text, rect, image, figure, embed, sketch, exercise). Three backends then
render the same elements:

  html   — a reveal.js deck for GitHub Pages (_site/<deck>/index.html)
  pptx   — an editable PowerPoint via python-pptx, post-processed by
           deckgen.classpoint (ait4x master, animations, ClassPoint)
  png    — a PIL preview of every slide + a text-overflow check

Paths — where the assets are read from and the outputs written — come from the
Project (see project.py), not from this file's location.

Design tokens come from the ait4x design system (scaffold/site/vendor/ait4x-colors_and_type.css).
Font names follow PPTX-EXPORT.md: the family name carries the weight.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from .project import current

HERE = Path(__file__).resolve().parent
FONT_DIR = HERE / 'fonts'      # Inter + JetBrains Mono ship with the package
JS_DIR = HERE / 'js'
SCAFFOLD_SITE = HERE / 'scaffold' / 'site'   # the vendor bundle (reveal.js, p5) staged into _site/
W, H = 1920, 1080
PX = 6350  # EMU per px on a 1920x1080 (13.333 x 7.5 in) slide
PT = 0.5   # points per design px on that slide (1920 px = 13.333 in = 960 pt)

# One size for code, everywhere. Code panels used to drift between 22, 24 and 26 px
# depending on which layout drew them; a student reading a projected slide should meet
# the same typeface at the same size every time code appears. 30 px = 15 pt in PowerPoint.
CODE = 30
CODE_SMALL = 26   # only where a panel genuinely cannot fit CODE (long lines, dense output)

# ───────────────────────── tokens (ait4x / PolyU Design) ─────────────────────────
INK, WHITE, PAPER, GRAY, TEAL = '#000B1C', '#FFFFFF', '#F4F4F2', '#BBBCB9', '#64C2C3'
TXT, MUTED, LINE, LINE_STRONG = '#2A323D', '#5C6470', '#E1E1DE', '#1F2832'
ORANGE, VIOLET, PINK, YELLOW, GREEN, BLUE, DEEP_TEAL, RED = '#ED6D24', '#943890', '#E94D7F', '#F6AD00', '#6FBA2C', '#146AB5', '#00544C', '#E42519'
MUTED_ON_INK, LIGHT_ON_INK = '#B3B7BE', '#D3E7E8'
# secondary matrix value scales used for bands
YELLOWS = ['#F6AD00', '#F0BD60', '#F4CE89', '#F8DEB1']
VIOLETS = ['#943890', '#A167A1', '#B68EBB', '#CDB5D4', '#E5DAEB', '#F2ECF5']
TEALS = ['#64C2C3', '#9FCED0', '#B9DBDC', '#D3E7E8', '#E9F3F4']
ORANGES = ['#ED6D24', '#E38E5D', '#EBAC83', '#F2C9AC', '#F9E5D6']
PINKS = ['#E94D7F', '#DF7F99', '#E7A3B4', '#EFC4CF', '#F7E3E8']

# font key -> (PowerPoint family, bold flag, css weight, PIL variation, css family, is_mono)
FONTS = {
    'black':    ('Inter Black', False, 900, 'Black', 'Inter', False),
    'xbold':    ('Inter ExtraBold', False, 800, 'ExtraBold', 'Inter', False),
    'bold':     ('Inter', True, 700, 'Bold', 'Inter', False),
    'semibold': ('Inter SemiBold', False, 600, 'SemiBold', 'Inter', False),
    'medium':   ('Inter Medium', False, 500, 'Medium', 'Inter', False),
    'body':     ('Inter', False, 400, 'Regular', 'Inter', False),
    'light':    ('Inter Light', False, 300, 'Light', 'Inter', False),
    'mono':     ('JetBrains Mono', False, 400, 'Regular', 'JetBrains Mono', True),
    'monomed':  ('JetBrains Mono Medium', False, 500, 'Medium', 'JetBrains Mono', True),
    'monobold': ('JetBrains Mono', True, 700, 'Bold', 'JetBrains Mono', True),
}


# ───────────────────────── element model ─────────────────────────
@dataclass
class Run:
    text: str
    font: str = 'body'
    size: int = 36
    color: str = INK
    spc: float = 0.0        # letter spacing in em
    caps: bool = False
    italic: bool = False
    url: str | None = None


@dataclass
class Para:
    runs: list
    align: str = 'l'        # l c r
    lh: float = 1.2         # line height, multiple of font size
    before: int = 0         # space before, px
    bullet: bool = False    # orange dot bullet (level 2 of the ait4x master)


@dataclass
class Text:
    x: int; y: int; w: int; h: int
    paras: list
    valign: str = 't'       # t m b
    name: str = ''
    kind: str = 'text'


@dataclass
class Rect:
    x: int; y: int; w: int; h: int
    fill: str | None = None
    stroke: str | None = None
    stroke_w: int = 2
    name: str = ''
    kind: str = 'rect'


@dataclass
class Image:
    x: int; y: int; w: int; h: int
    src: str                # path relative to deck/assets or absolute
    fit: str = 'cover'      # cover | contain
    name: str = ''
    placeholder: str = ''   # drawn instead when the file is missing (the still of a sketch not made yet)
    kind: str = 'image'


@dataclass
class Figure:
    """A drawn illustration: svg for html, png for pptx/preview."""
    x: int; y: int; w: int; h: int
    svg: str
    png: str                # path to a rendered png
    name: str = ''
    kind: str = 'figure'


@dataclass
class Embed:
    """YouTube embed. html: iframe; pptx/png: thumbnail + link."""
    x: int; y: int; w: int; h: int
    yt: str
    thumb: str | None = None
    name: str = ''
    kind: str = 'embed'


@dataclass
class Sketch:
    """A live, interactive p5.js sketch, html only: an iframe laid over the twin that stands in
    for it in the pptx and the PDF (a drawn figure, or the still in deck/assets/sketches made by
    `deckgen snap`). p5 ships in the vendor bundle, so the deck also runs offline. The page scales
    the canvas to the frame (mouseX/mouseY stay right), forwards the deck's navigation keys to
    reveal.js, restarts on R, and shows a title bar when opened on its own."""
    x: int; y: int; w: int; h: int
    name: str               # file name of the page (<deck>/sketches/<name>.html) and of the still
    code: str               # the sketch, as the students see it on the slide
    cw: int = 600           # the sketch's own canvas size; the page scales it to the frame
    ch: int = 600
    hint: str = ''          # what to do with it, shown in the LIVE chip: 'move the mouse · click to reseed'
    extra: str = ''         # js appended after the code (interaction the slide's code panel does not show)
    sound: bool = False     # also load p5.sound (oscillators, the microphone, playback)
    twin: bool = True       # False when a drawn figure stands in for it (no still needed)
    kind: str = 'sketch'


def twin_png(name):
    """Where `deckgen snap` puts the still of a sketch (the pptx / PDF stand-in): deck/assets/sketches/<name>.png."""
    return current().sketches / f'{name}.png'


@dataclass
class Exercise:
    """A runnable Python drill.

    html: an editable code box, a Run button, an output pane and a pass/fail check,
    executed in the browser by Pyodide — no install, works on a student's laptop
    during class. pptx and png: the same code as a static panel, because PowerPoint
    cannot run Python and a printed slide should still show the exercise.

    `check` is Python run after the student's code, in the same namespace, with the
    captured stdout bound to `_out`. It passes by setting `ok = True`; anything it
    puts in `msg` is shown to the student. `expect` is the shorthand for the common
    case: stdout must equal this string once both sides are stripped.
    """
    x: int; y: int; w: int; h: int
    code: str                       # starter code, shown in the editor
    check: str = ''                 # python; set ok = True to pass
    expect: str | None = None       # or: expected stdout, compared stripped
    eid: str = ''                   # stable id — keys the saved answer in localStorage
    label: str = ''                 # 'EXERCISE 01'
    hint: str = ''
    rows: int = 8                   # editor height, in lines
    size: int = CODE                # code size: CODE, or stepped down by layouts.exercise so the code fits the panel
    name: str = ''
    kind: str = 'exercise'


@dataclass
class Slide:
    bg: str = WHITE
    els: list = field(default_factory=list)
    notes: str = ''
    cp: dict | None = None      # ClassPoint activity for build.py
    report: str = ''            # after class: the public activity page, set by attach_reports
    title: str = ''             # for the outline / html title
    html_only: list = field(default_factory=list)  # elements only for html (e.g. live demos)


def exercise_static(el):
    """The pptx/png rendering of an Exercise: the panel and the code, no controls."""
    from .syntax import code_runs
    body = [Para(rs, 'l', 1.45) for rs in code_runs(el.code, el.size)]
    out = [Rect(el.x, el.y, el.w, el.h, PAPER, name=el.name or 'exercise-panel'),
           Text(el.x + 40, el.y + 32, el.w - 80, el.h - 64, body, 't', name='code')]
    if el.label:
        out.insert(1, Text(el.x + 40, el.y - 44, el.w - 80, 36,
                           [Para(runs(el.label, 'monomed', 22, ORANGE, spc=0.14, caps=True), 'l', 1.2)], 't'))
    return out


# ───────────────────────── inline markup ─────────────────────────
# **bold**  ·  [text](url)  ·  {orange:text} {teal:text} {muted:text} {mono:text}
_TOKEN = re.compile(r'(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|\{[a-z]+:[^}]+\})')
# Each token kind, anchored, for re-identifying a part after the split. Matching them by
# startswith() instead used to mangle any line that merely began with '{' and held a
# colon — a dict literal at the start of a line of code — into its own truncated tail.
_BOLD = re.compile(r'\*\*(.+?)\*\*\Z', re.S)
_LINK = re.compile(r'\[([^\]]+)\]\(([^)]+)\)\Z')
_TAG = re.compile(r'\{([a-z]+):([^}]+)\}\Z')
_COLORS = {'orange': ORANGE, 'teal': TEAL, 'muted': MUTED, 'violet': VIOLET, 'pink': PINK,
           'white': WHITE, 'ink': INK, 'green': GREEN, 'yellow': YELLOW, 'blue': BLUE}


def runs(text, font='body', size=36, color=INK, spc=0.0, caps=False, bold_font=None, italic=False):
    """Parse light inline markup into Run objects."""
    bold_font = bold_font or ('bold' if font in ('body', 'light', 'medium') else font)
    out = []
    for part in _TOKEN.split(text):
        if not part:
            continue
        bold, link, tag = _BOLD.match(part), _LINK.match(part), _TAG.match(part)
        if bold:
            out.append(Run(bold.group(1), bold_font, size, color, spc, caps, italic))
        elif link:
            out.append(Run(link.group(1), font, size, color, spc, caps, italic, url=link.group(2)))
        elif tag:
            key, val = tag.group(1), tag.group(2)
            if key == 'mono':
                out.append(Run(val, 'mono', size, color, spc, caps, italic))
            elif key == 'bold':
                out.append(Run(val, bold_font, size, color, spc, caps, italic))
            else:
                out.append(Run(val, font, size, _COLORS.get(key, color), spc, caps, italic))
        else:
            out.append(Run(part, font, size, color, spc, caps, italic))
    return out


def T(x, y, w, h, text, font='body', size=36, color=INK, lh=1.2, align='l', valign='t', spc=0.0, caps=False, name='', italic=False):
    """Single-paragraph text box; '\n' starts a new paragraph."""
    paras = [Para(runs(line, font, size, color, spc, caps, italic=italic), align, lh) for line in text.split('\n')]
    return Text(x, y, w, h, paras, valign, name)


def P(text, font='body', size=36, color=INK, lh=1.4, align='l', before=0, bullet=False, spc=0.0, caps=False):
    return Para(runs(text, font, size, color, spc, caps), align, lh, before, bullet)


def eyebrow(x, y, text, color=MUTED, w=1500, size=24, name=''):
    return T(x, y, w, 40, text, 'monomed', size, color, lh=1.2, spc=0.14, caps=True, name=name)


# ───────────────────────── measuring (PIL) ─────────────────────────
_font_cache = {}


def pil_font(key, size):
    from PIL import ImageFont
    fam, _b, _w, var, _css, mono = FONTS[key]
    path = FONT_DIR / ('JetBrainsMono-Variable.ttf' if mono else 'Inter-Variable.ttf')
    k = (key, size)
    if k not in _font_cache:
        f = ImageFont.truetype(str(path), size)
        try:
            f.set_variation_by_name(var)
        except Exception:
            names = [n.decode() if isinstance(n, bytes) else n for n in f.get_variation_names()]
            if var in names:
                f.set_variation_by_name(var)
        _font_cache[k] = f
    return _font_cache[k]


def natural_lh(key, size):
    """Ascent+descent of the font at size, in px (what PowerPoint calls single spacing)."""
    a, d = pil_font(key, size).getmetrics()
    return a + d


def run_width(run):
    txt = run.text.upper() if run.caps else run.text
    f = pil_font(run.font, run.size)
    return f.getlength(txt) + run.spc * run.size * len(txt)


def wrap_para(para, width):
    """Greedy wrap over runs; returns list of lines, each a list of (Run, text) fragments."""
    lines, cur, cur_w = [], [], 0.0
    for run in para.runs:
        words = re.split(r'(\s+)', run.text)
        for wd in words:
            if wd == '':
                continue
            if '\n' in wd:
                wd = wd.replace('\n', ' ')
            frag = Run(wd, run.font, run.size, run.color, run.spc, run.caps, run.italic, run.url)
            fw = run_width(frag)
            if cur and cur_w + fw > width and wd.strip():
                lines.append(cur); cur, cur_w = [], 0.0
                if not wd.strip():
                    continue
            if not cur and not wd.strip() and lines:
                continue          # drop the whitespace a wrap lands on; keep a paragraph's own indentation (code)
            cur.append((frag, wd)); cur_w += fw
    if cur:
        lines.append(cur)
    return lines or [[]]


def text_height(el):
    total = 0.0
    for p in el.paras:
        size = max((r.size for r in p.runs), default=24)
        n = len(wrap_para(p, el.w))
        total += p.before + n * p.lh * size
    return total


# ───────────────────────── HTML backend ─────────────────────────
CSS = """
@font-face{font-family:'Inter';src:url('../vendor/fonts/Inter-Variable.ttf') format('truetype');font-weight:100 900;font-display:swap}
@font-face{font-family:'JetBrains Mono';src:url('../vendor/fonts/JetBrainsMono-Variable.ttf') format('truetype');font-weight:100 800;font-display:swap}
:root{--font-d:'Inter','Helvetica Now','Helvetica Neue',Helvetica,Arial,sans-serif;--font-m:'JetBrains Mono',Menlo,Consolas,monospace}
html,body{background:#000B1C}
.reveal{font-family:var(--font-d);color:#000B1C;-webkit-font-smoothing:antialiased;font-feature-settings:"ss01","cv11"}
.reveal .slides{text-align:left}
.reveal .slides section{padding:0;width:1920px;height:1080px;top:0;left:0;position:absolute;overflow:hidden}
.el{position:absolute;box-sizing:border-box;margin:0;display:flex;flex-direction:column;overflow:visible}
.el p{margin:0;white-space:pre-wrap;overflow-wrap:break-word}
.el p.bullet{padding-left:48px;text-indent:-48px}
.el p.bullet::before{content:'·';color:#ED6D24;font-family:var(--font-d);font-weight:900;display:inline-block;width:48px;text-indent:0}
.el a{color:inherit;text-decoration:underline;text-decoration-thickness:2px;text-underline-offset:.12em}
.el a:hover{color:#00544C}
.img{position:absolute;overflow:hidden}
.img img{width:100%;height:100%;display:block}
.fig{position:absolute}
.fig svg{width:100%;height:100%;display:block}
.embed{position:absolute;background:#000}
.embed iframe{width:100%;height:100%;border:0}
.cp{position:absolute;left:1450px;top:908px;width:350px;height:92px;border:2px dashed #ED6D24;color:#ED6D24;font:500 22px/1 var(--font-m);letter-spacing:.14em;text-transform:uppercase;display:flex;align-items:center;justify-content:center;gap:12px}
.cp b{width:12px;height:12px;border-radius:50%;background:#ED6D24;display:inline-block}
a.cp{position:absolute;border-style:solid;text-decoration:none;cursor:pointer;left:auto;right:120px;width:auto;padding:0 28px;white-space:nowrap}a.cp:hover{background:#ED6D24;color:#fff}a.cp:hover b{background:#fff}
.embed .print-only{display:none}
.embed.sketch{background:#fff}
.embed.sketch .live{position:absolute;right:0;bottom:0;max-width:100%;box-sizing:border-box;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;background:#000B1C;color:#fff;font:500 17px/1 var(--font-m);letter-spacing:.12em;text-transform:uppercase;padding:10px 14px;display:flex;align-items:center;gap:12px;pointer-events:none;opacity:.92;transition:opacity .5s}
.embed.sketch .live.dim{opacity:.14}
.embed.sketch .live b{width:10px;height:10px;border-radius:50%;background:#ED6D24;display:inline-block;animation:pulse 1.6s infinite}
.embed.sketch .live a{color:#fff;text-decoration:none;pointer-events:auto;padding:0 4px;font-size:22px}
@keyframes pulse{50%{opacity:.35}}
@media print{.cp{display:none}.embed iframe{display:none}.embed.sketch{display:none!important}.embed .print-only{display:block;width:100%;height:100%;object-fit:cover}}
.reveal .progress{height:4px;color:#ED6D24}
.reveal .backgrounds{background:#000B1C}

/* ── runnable exercises (vendor/pyodide-console.js) ── */
.ex{position:absolute;display:flex;flex-direction:column;background:#F4F4F2;border:2px solid #E1E1DE;box-sizing:border-box}
.ex-head{display:flex;justify-content:space-between;align-items:baseline;padding:14px 20px 0}
.ex-label{font:500 22px/1 var(--font-m);letter-spacing:.14em;text-transform:uppercase;color:#ED6D24}
.ex-status{font:500 22px/1 var(--font-m);letter-spacing:.08em;text-transform:uppercase;color:#5C6470}
.ex.pass .ex-status{color:#00544C} .ex.fail .ex-status{color:#ED6D24} .ex.busy .ex-status{color:#5C6470}
.ex.pass{border-color:#64C2C3} .ex.fail{border-color:#ED6D24}
/* the editor: a transparent textarea over a coloured <pre> with the same metrics, so the
   caret and the selection are the browser's own and the colours are ours */
.ex-editor{position:relative;flex:1 1 auto;min-height:0;margin:12px 20px 0;background:#fff}
.ex-editor:focus-within{box-shadow:inset 0 0 0 2px #64C2C3}
.ex-hl,.ex-code{position:absolute;inset:0;box-sizing:border-box;margin:0;padding:16px;border:0;
  font:400 1em/1.45 var(--font-m);tab-size:4;white-space:pre;overflow:auto;color:#000B1C}
.ex-hl{pointer-events:none;overflow:hidden;background:transparent}
.ex-code{background:transparent;color:transparent;caret-color:#000B1C;resize:none;outline:0}
.ex-bar{display:flex;align-items:center;gap:12px;padding:12px 20px}
.ex-bar button{font:500 24px/1 var(--font-m);letter-spacing:.08em;text-transform:uppercase;padding:10px 22px;
  border:0;cursor:pointer;background:#000B1C;color:#fff}
.ex-bar button.ex-reset,.ex-bar button.ex-send{background:transparent;color:#5C6470;padding:10px 8px}
.ex-bar button:hover{background:#ED6D24;color:#fff}
.ex-hint{font:400 22px/1.3 var(--font-d);color:#5C6470;margin-left:auto;text-align:right}
.ex-out{margin:0 20px 20px;padding:0;max-height:34%;overflow:auto;white-space:pre-wrap;
  font:400 __CODE_SMALL__px/1.4 var(--font-m);color:#2A323D}
.ex-out:empty{display:none}
.ex.ex-fig .ex-out{max-height:none;overflow:hidden;text-align:center}
.ex-figure svg{display:block;margin:0 auto;height:430px;max-width:100%;background:#fff;overflow:visible;padding:6px;box-sizing:border-box}
.ex-hl,.ex-code,#pyc-log,#pyc-in{font-variant-ligatures:none}
.ex-out .err{color:#E42519}
.ex-out .msg{color:#00544C}

/* the global console: backtick, or the button bottom-left */
#pyc{position:fixed;left:0;right:0;bottom:0;height:44vh;background:#000B1C;color:#D3E7E8;z-index:60;
  display:none;flex-direction:column;box-shadow:0 -8px 40px rgba(0,0,0,.5)}
#pyc.on{display:flex}
#pyc .pyc-log{flex:1 1 auto;overflow:auto;padding:20px 28px;margin:0;white-space:pre-wrap;font:400 24px/1.5 var(--font-m)}
#pyc .pyc-log .in{color:#64C2C3} #pyc .pyc-log .err{color:#ED6D24}
#pyc .pyc-in{border:0;outline:0;resize:none;box-sizing:border-box;width:100%;flex:0 0 auto;
  background:#0A1526;color:#fff;padding:18px 28px;font:400 26px/1.4 var(--font-m);
  tab-size:4;white-space:pre;overflow:auto;max-height:22vh}
#pyc .pyc-bar{display:flex;justify-content:space-between;align-items:center;padding:10px 28px;
  font:500 20px/1 var(--font-m);letter-spacing:.14em;text-transform:uppercase;color:#5C6470;border-bottom:1px solid #1F2832}
#pyc .pyc-bar button{background:none;border:0;color:#5C6470;font:inherit;cursor:pointer}
#pyc .pyc-bar button:hover{color:#ED6D24}
#pyc-open{position:fixed;left:16px;bottom:16px;z-index:59;background:#000B1C;color:#fff;border:0;cursor:pointer;
  font:500 20px/1 var(--font-m);letter-spacing:.14em;text-transform:uppercase;padding:12px 18px;opacity:.55}
#pyc-open:hover{opacity:1;background:#ED6D24}
#pyload{position:fixed;left:50%;top:50%;transform:translate(-50%,-50%);z-index:70;background:#000B1C;color:#fff;
  padding:28px 36px;font:500 24px/1.4 var(--font-m);letter-spacing:.08em;display:none}
#pyload.on{display:block}

@media print{.ex-bar,.ex-status,.ex-out,#pyc,#pyc-open,#pyload,.ex-code{display:none!important}
  .ex-hl{overflow:visible}}

/* ── reading view: the same slides, reflowed, for phones ── */
.handout{display:none;background:#fff;color:#000B1C;max-width:44rem;margin:0 auto;padding:0 20px 96px;
  font-family:var(--font-d);-webkit-text-size-adjust:100%}
body.ho .handout{display:block}
body.ho .reveal{display:none}
body.ho{overflow:auto;height:auto;background:#fff}
.ho-bar{position:sticky;top:0;z-index:20;background:#fff;border-bottom:1px solid #E1E1DE;
  display:flex;justify-content:space-between;align-items:center;gap:16px;padding:14px 0;margin-bottom:8px}
.ho-title{font:800 17px/1.2 var(--font-d);letter-spacing:-.02em;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.ho-bar button,#ho-open{font:500 12px/1 var(--font-m);letter-spacing:.14em;text-transform:uppercase;
  background:#000B1C;color:#fff;border:0;padding:11px 14px;cursor:pointer;flex:none;border-radius:2px}
#ho-open{position:fixed;right:12px;bottom:12px;z-index:59;opacity:.85;display:none}
.ho-slide{position:relative;padding:28px 0 32px;border-bottom:1px solid #E1E1DE}
.ho-n{position:absolute;right:0;top:30px;font:500 11px/1 var(--font-m);color:#BBBCB9;letter-spacing:.14em}
.ho-slide h2{font:800 30px/1.1 var(--font-d);letter-spacing:-.03em;margin:0 0 14px}
.ho-slide h3{font:800 22px/1.2 var(--font-d);letter-spacing:-.02em;margin:18px 0 10px}
.ho-slide p{font:400 17px/1.55 var(--font-d);color:#2A323D;margin:0 0 12px}
.ho-slide ul{margin:0 0 14px;padding-left:20px} .ho-slide li{font:400 17px/1.55 var(--font-d);color:#2A323D;margin:0 0 7px}
.ho-slide p.ho-eyebrow{font:500 11px/1.4 var(--font-m);letter-spacing:.16em;color:#5C6470;margin:0 0 10px}
.ho-slide p.ho-label{font:500 13px/1.4 var(--font-m);letter-spacing:.06em;color:#ED6D24;margin:0 0 6px}
.ho-slide p.ho-choice{display:flex;gap:10px;margin:0 0 8px}
.ho-slide p.ho-choice b{flex:none;width:26px;height:26px;background:#000B1C;color:#fff;
  font:500 13px/26px var(--font-m);text-align:center}
.ho-slide code{font-family:var(--font-m);font-size:.9em;background:#F4F4F2;padding:1px 4px}
.ho-slide strong{font-weight:700}
.ho-slide a{color:#00544C;text-decoration:underline;text-underline-offset:.14em}
.ho-slide img{width:100%;height:auto;display:block;margin:14px 0}
.ho-fig svg{width:100%;height:auto;display:block;margin:14px 0}
.ho-code{background:#F4F4F2;padding:14px;margin:0 0 14px;overflow-x:auto;
  font:400 14px/1.5 var(--font-m);white-space:pre;-webkit-overflow-scrolling:touch}
.ho-cp{font:500 11px/1.4 var(--font-m);letter-spacing:.14em;text-transform:uppercase;color:#ED6D24}
/* the moved exercise: inline styles carry the deck geometry, so they need overriding */
.ho-ex .ex{position:static!important;left:auto!important;top:auto!important;
  width:100%!important;height:auto!important;margin:14px 0}
.ho-ex .ex-head{padding:12px 14px 0} .ho-ex .ex-label,.ho-ex .ex-status{font-size:11px}
.ho-ex .ex-editor{margin:10px 14px 0;min-height:9.5em;height:auto}
.ho-ex .ex-editor{font-size:16px} .ho-ex .ex-hl,.ho-ex .ex-code{padding:12px;line-height:1.5}
.ho-ex .ex-hl{position:static;visibility:hidden;overflow:visible}  /* sizes the box to the code; the absolute textarea paints over it */
.ho-ex .ex-code{color:#000B1C;background:#fff}
.ho-ex .ex-bar{padding:10px 14px;flex-wrap:wrap}
.ho-ex .ex-bar button{font-size:12px;padding:11px 16px}
.ho-ex .ex-hint{font-size:12px;margin-left:auto}
.ho-ex .ex-out{margin:0 14px 14px;font-size:14px;max-height:none}
.ho-ex .ex-figure svg{height:auto;width:100%}
@media (max-width:900px){#ho-open{display:block}}
/* already reading: the way out is the Deck view button in the sticky bar, and leaving
   this one up collides with the console button in the opposite corner */
body.ho #ho-open{display:none!important}
@media print{.handout,#ho-open{display:none!important} body.ho .reveal{display:block!important}}
"""

# Injected only when a deck actually has exercises or asks for the console — a deck of
# pure prose should not make 112 laptops fetch a Python runtime.
PYODIDE_HTML = '''
<button id="pyc-open" type="button" title="Python console (`)">Python ›_</button>
<div id="pyload">Loading Python…</div>
<div id="pyc"><div class="pyc-bar"><span>Python console — enter runs, shift+enter adds a line, paste keeps its indentation, ` to close</span><button id="pyc-clear" type="button">clear</button></div>
<pre class="pyc-log"></pre><textarea class="pyc-in" rows="1" spellcheck="false" autocapitalize="off" autocorrect="off" placeholder=">>>"></textarea></div>
<script>window.PYODIDE_URL={pyodide_url!r};</script>
<script src="../vendor/pyodide-console.js"></script>
'''

HTML_TMPL = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<link rel="stylesheet" href="../vendor/reveal/reset.css">
<link rel="stylesheet" href="../vendor/reveal/reveal.css">
<style>{css}</style>
</head>
<body>
<div class="reveal"><div class="slides">
{slides}
</div></div>
<main class="handout" hidden>
<div class="ho-bar"><span class="ho-title">{title}</span><button id="ho-deck" type="button">Deck view</button></div>
{handout}
</main>
<button id="ho-open" type="button" title="Reading view">Reading view</button>
<script src="../vendor/reveal/reveal.js"></script>
<script src="../vendor/reveal/plugin/notes/notes.js"></script>
{pyodide}
<script src="../vendor/handout.js"></script>
<script>
Reveal.initialize({{width:1920,height:1080,margin:0,minScale:0.05,maxScale:4,center:false,hash:true,transition:'none',
  backgroundTransition:'none',controls:false,progress:true,slideNumber:false,plugins:[RevealNotes],
  pdfMaxPagesPerSlide:1,pdfSeparateFragments:false,keyboard:{{}},
  // reveal's own scroll view activates below 435px and CLONES every slide, which on a
  // phone gave two of every exercise — duplicate ids, and Run doing nothing on whichever
  // copy the browser resolved first. The reading view is our answer to small screens.
  scrollActivationWidth:null}});
// the LIVE chip on a sketch shows its hint for four seconds, then fades so it does not cover the picture
var dimTimer = null;
function chips() {{
  clearTimeout(dimTimer);
  document.querySelectorAll('.live').forEach(function (c) {{ c.classList.remove('dim'); }});
  dimTimer = setTimeout(function () {{ document.querySelectorAll('.present .live').forEach(function (c) {{ c.classList.add('dim'); }}); }}, 4000);
}}
Reveal.on('ready', chips); Reveal.on('slidechanged', chips);
// a live sketch that has the keyboard focus relays the deck's navigation keys (SKETCH_TMPL in deckgen.core)
addEventListener('message', function (e) {{
  var d = e.data; if (!d || d.deckgen !== 'key') return;
  if (document.activeElement && document.activeElement.blur) document.activeElement.blur();   // the deck takes the keyboard back
  window.focus();
  var k = d.key;
  if (k === 'ArrowRight' || k === 'ArrowDown' || k === 'PageDown' || (k === ' ' && !d.shift)) Reveal.next();
  else if (k === 'ArrowLeft' || k === 'ArrowUp' || k === 'PageUp' || (k === ' ' && d.shift)) Reveal.prev();
  else if (k === 'Home') Reveal.slide(0);
  else if (k === 'End') Reveal.slide(Reveal.getTotalSlides() - 1);
  else if (k === 'Escape' || k === 'o' || k === 'O') Reveal.toggleOverview();
  else if (k === 's' || k === 'S') Reveal.getPlugin('notes').open();
  else if (k === 'f' || k === 'F') {{ document.fullscreenElement ? document.exitFullscreen() : document.documentElement.requestFullscreen(); }}
}});
</script>
</body>
</html>
"""


# The standalone page of a live sketch (<deck>/sketches/<name>.html), also what the deck's iframe loads.
# @@ tokens, not str.format: the sketch code is full of braces.
SKETCH_TMPL = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>@@NAME@@ · @@COURSE@@ · p5.js</title>
<style>
html,body{margin:0;height:100%;background:#fff;overflow:hidden;font-family:'JetBrains Mono',Menlo,Consolas,monospace}
canvas{position:absolute;left:50%;top:50%;transform:translate(-50%,-50%);display:block}
body.top{background:#F4F4F2}
#bar{display:none;position:fixed;left:0;right:0;top:0;height:44px;background:#000B1C;color:#fff;font-size:12px;letter-spacing:.14em;text-transform:uppercase;align-items:center;padding:0 16px;gap:16px;z-index:9;white-space:nowrap;overflow:hidden}
#bar b{color:#ED6D24;font-weight:500}#bar span{opacity:.6}#bar a{color:#fff;margin-left:auto;text-decoration:none;opacity:.6}
body.top #bar{display:flex}body.top canvas{top:calc(50% + 22px)}
</style>
<script src="../../vendor/p5/p5.min.js"></script>
@@SOUND@@
</head><body>
<div id="bar"><b>@@COURSE@@</b> @@NAME@@ <span>@@HINT@@ · R restarts</span><a href="../">← the deck</a></div>
<script>
@@CODE@@
@@EXTRA@@
</script>
<script>
(function () {
  var top = (self === window.top) && !/[?&]snap/.test(location.search);   // on its own: a title bar; ?snap: none (deckgen snap)
  if (top) document.body.classList.add('top');
  var CW = @@CW@@, CH = @@CH@@;
  function fit() {                                   // scale by css size, not transform: p5 maps mouseX from scrollWidth
    var c = document.querySelector('canvas'); if (!c) return;
    var pad = top ? 44 : 0, s = Math.min(innerWidth / CW, (innerHeight - pad) / CH);
    var w = Math.round(CW * s) + 'px', h = Math.round(CH * s) + 'px';
    if (c.style.width !== w) c.style.width = w;
    if (c.style.height !== h) c.style.height = h;
  }
  new MutationObserver(fit).observe(document.documentElement, {childList: true, subtree: true, attributes: true, attributeFilter: ['width', 'height', 'style']});
  addEventListener('resize', fit); fit();
  var NAV = {ArrowLeft: 1, ArrowRight: 1, ArrowUp: 1, ArrowDown: 1, PageUp: 1, PageDown: 1, Home: 1, End: 1, Escape: 1, ' ': 1, s: 1, S: 1, o: 1, O: 1, f: 1, F: 1};
  addEventListener('keydown', function (e) {
    if (e.defaultPrevented || e.metaKey || e.ctrlKey || e.altKey) return;   // the sketch used the key (keyPressed returned false)
    if (e.key === 'r' || e.key === 'R') { location.reload(); return; }
    if (NAV[e.key] && !top) { parent.postMessage({deckgen: 'key', key: e.key, shift: e.shiftKey}, '*'); e.preventDefault(); }
  });
})();
</script>
</body></html>
"""


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def attr(s):
    """esc() is for text nodes and leaves quotes alone — inside an attribute a bare
    quote ends it early. Python checks are full of them."""
    return esc(s).replace('"', '&quot;').replace("'", '&#39;')


def css_font(run):
    fam, bold, weight, _v, css_fam, mono = FONTS[run.font]
    style = f"font-family:'{fam}',{'var(--font-m)' if mono else 'var(--font-d)'};font-weight:{weight};font-size:{run.size}px;color:{run.color};"
    if run.spc:
        style += f'letter-spacing:{run.spc}em;'
    if run.caps:
        style += 'text-transform:uppercase;'
    if run.italic:
        style += 'font-style:italic;'
    return style


def html_text(el):
    valign = {'t': 'flex-start', 'm': 'center', 'b': 'flex-end'}[el.valign]
    out = [f'<div class="el" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px;justify-content:{valign}">']
    for p in el.paras:
        size = max((r.size for r in p.runs), default=24)
        align = {'l': 'left', 'c': 'center', 'r': 'right'}[p.align]
        cls = ' class="bullet"' if p.bullet else ''
        out.append(f'<p{cls} style="text-align:{align};line-height:{p.lh};margin-top:{p.before}px;font-size:{size}px">')
        for r in p.runs:
            t = esc(r.text)
            span = f'<span style="{css_font(r)}">{t}</span>'
            if r.url:
                span = f'<a href="{attr(r.url)}" target="_blank" rel="noopener">{span}</a>'
            out.append(span)
        out.append('</p>')
    out.append('</div>')
    return ''.join(out)


def html_slide(s, i, out_dir, assets_out, assets_rel):
    bg = f' data-background-color="{s.bg}"'
    parts = [f'<section{bg} data-slide="{i}">']
    for el in s.els + s.html_only:
        if el.kind == 'text':
            parts.append(html_text(el))
        elif el.kind == 'rect':
            st = f'left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px;'
            if el.fill:
                st += f'background:{el.fill};'
            if el.stroke:
                st += f'border:{el.stroke_w}px solid {el.stroke};'
            parts.append(f'<div class="el" style="{st}"></div>')
        elif el.kind == 'image':
            src = copy_asset(el.src, assets_out)
            fit = 'cover' if el.fit == 'cover' else 'contain'
            if src is None:   # the still of a sketch not made yet: the live frame covers the spot on screen, print shows the label
                parts.append(f'<div class="el" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px;background:{PAPER};align-items:center;justify-content:center;font:500 22px/1 var(--font-m);color:{MUTED};letter-spacing:.14em;text-transform:uppercase">{esc(el.placeholder or el.src)}</div>')
            else:
                parts.append(f'<div class="img" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px"><img src="{assets_rel}/{src}" alt="" style="object-fit:{fit}"></div>')
        elif el.kind == 'figure':
            parts.append(f'<div class="fig" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px">{el.svg}</div>')
        elif el.kind == 'exercise':
            parts.append(html_exercise(el))
        elif el.kind == 'sketch':
            write_sketch_page(el, out_dir)
            hint = ' · ' + esc(el.hint) if el.hint else ''
            chip = f'<div class="live"><b></b>Live{hint}<a href="sketches/{el.name}.html" target="_blank" rel="noopener" title="open the sketch in its own tab">↗</a></div>'
            parts.append(f'<div class="embed sketch" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px"><iframe data-src="sketches/{el.name}.html" title="{attr(el.name)}"></iframe>{chip}</div>')
        elif el.kind == 'embed':
            thumb = f'<img class="print-only" src="{assets_rel}/{copy_asset(el.thumb, assets_out)}" alt="">' if el.thumb else ''  # the pdf shows the thumbnail
            parts.append(f'<div class="embed" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px"><iframe data-src="https://www.youtube-nocookie.com/embed/{el.yt}?rel=0" allow="accelerometer; autoplay; clipboard-write; encrypted-media; gyroscope; picture-in-picture" allowfullscreen loading="lazy"></iframe>{thumb}</div>')
    if s.cp:
        label = {'word_cloud': 'Word cloud', 'multiple_choice': 'Multiple choice', 'short_answer': 'Short answer', 'image_upload': 'Image upload'}[s.cp['type']]
        # Before the class the chip says what kind of activity this is. After it, the
        # same chip is the way to the answers — a link, not a badge that looks like one.
        if s.report:
            parts.append(f'<a class="cp" data-classpoint href="{attr(s.report)}" target="_blank" rel="noopener"><b></b>{label} · see the answers</a>')
        else:
            parts.append(f'<div class="cp" data-classpoint><b></b>{label}</div>')
    if s.notes:
        parts.append(f'<aside class="notes">{esc(s.notes)}</aside>')
    parts.append('</section>')
    return '\n'.join(parts)


def html_exercise(el):
    """Editor + Run + output. The check travels as a data attribute; the runtime in
    vendor/pyodide-console.js does the executing."""
    from .syntax import html as syntax_html
    eid = attr(el.eid or el.name or 'ex')
    attrs = f'class="ex" id="ex-{eid}" data-eid="{eid}" data-rows="{el.rows}"'
    if el.check:
        attrs += f' data-check="{attr(el.check)}"'
    if el.expect is not None:
        attrs += f' data-expect="{attr(el.expect)}"'
    head = f'<div class="ex-head"><span class="ex-label">{esc(el.label)}</span><span class="ex-status" aria-live="polite"></span></div>' if el.label else ''
    hint = f'<span class="ex-hint">{esc(el.hint)}</span>' if el.hint else ''
    return (f'<div {attrs} style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px">'
            f'{head}'
            f'<div class="ex-editor" style="font-size:{el.size}px"><pre class="ex-hl" aria-hidden="true">{syntax_html(el.code)}\n</pre>'
            f'<textarea class="ex-code" spellcheck="false" autocapitalize="off" autocorrect="off">{esc(el.code)}</textarea></div>'
            f'<div class="ex-bar"><button class="ex-run" type="button">Run</button>'
            f'<button class="ex-reset" type="button">Reset</button>'
            f'<button class="ex-send" type="button" title="Load this code into the Python console">&rsaquo;_</button>'
            f'{hint}</div>'
            f'<pre class="ex-out"></pre></div>')


# ───────────────────────── handout (the phone view) ─────────────────────────
# A 1920x1080 canvas of absolutely-positioned boxes cannot be made readable on a phone
# by scaling: reveal fits the slide to the viewport, so on a 390px screen 36px body text
# lands at about 7px and the exercise box is too small to type in. The content has to
# reflow, which means a second, semantic rendering of the same slide data.
#
# Roles are inferred from what the layouts already encode — size, font, caps, bullets —
# rather than by tagging every layout function. It is a heuristic, and it is allowed to
# be: this view is for reading and for doing the exercises, not for reproducing the deck.
# A short mono run that labels the thing after it: a choice letter, an agenda number,
# a timeline year, a start time. On the slide it sits beside its text; in a reflowed
# column it should join it rather than become an orphan line.
LABEL_PREFIX = re.compile(r'^(?:[A-F]|\d{1,4}|\d{1,2}:\d{2})$')


def _plain(el):
    return ''.join(r.text for p in el.paras for r in p.runs).strip()


def _role(el):
    runs_all = [r for p in el.paras for r in p.runs]
    if not runs_all:
        return 'skip'
    size = max(r.size for r in runs_all)
    mono = all(FONTS[r.font][5] for r in runs_all)
    # Code is declared by the layout, not inferred. Guessing it from "monospaced and
    # biggish" turned agenda numbers, timeline years and quote attributions into <pre>.
    if el.name == 'code':
        return 'code'
    if mono and all(r.caps for r in runs_all):
        return 'eyebrow'
    if mono:
        return 'label'
    if size >= 60:
        return 'h2'
    if size >= 40:
        return 'h3'
    return 'p'


def handout_runs(para, role='p'):
    # A heading is already bold and an eyebrow is already monospaced by the stylesheet;
    # re-marking their runs would nest <strong> inside <h2> and <code> inside an eyebrow.
    plain = role in ('h2', 'h3', 'eyebrow')
    out = []
    for r in para.runs:
        t = esc(r.text.upper() if r.caps else r.text)
        if plain:
            pass
        elif FONTS[r.font][5]:
            t = f'<code>{t}</code>'
        elif FONTS[r.font][2] >= 600:
            t = f'<strong>{t}</strong>'
        if r.color in (ORANGE, TEAL, VIOLET, PINK, GREEN, BLUE, DEEP_TEAL):
            t = f'<em style="color:{r.color};font-style:normal">{t}</em>'
        if r.url:
            t = f'<a href="{attr(r.url)}">{t}</a>'
        out.append(t)
    return ''.join(out) or '&nbsp;'


def handout_text(el, role):
    if role == 'code':
        def span(r):
            t = esc(r.text)
            return f'<span style="color:{r.color}">{t}</span>' if r.color not in (INK, LIGHT_ON_INK) else t
        return '<pre class="ho-code">' + '\n'.join(
            ''.join(span(r) for r in p.runs) for p in el.paras) + '</pre>'
    out, bullets = [], []

    def flush():
        if bullets:
            out.append('<ul>' + ''.join(f'<li>{b}</li>' for b in bullets) + '</ul>')
            bullets.clear()

    for para in el.paras:
        body = handout_runs(para, role)
        if body == '&nbsp;':
            continue
        if para.bullet:
            bullets.append(body)
            continue
        flush()
        tag = {'eyebrow': 'p class="ho-eyebrow"', 'label': 'p class="ho-label"',
               'h2': 'h2', 'h3': 'h3'}.get(role, 'p')
        out.append(f'<{tag}>{body}</{tag.split()[0]}>')
    flush()
    return ''.join(out)


def handout_slide(s, i, assets_out, assets_rel):
    parts = [f'<article class="ho-slide" id="ho-{i}"><span class="ho-n">{i}</span>']
    pending = ''          # a bare choice letter waits for the choice it labels
    for el in s.els + s.html_only:
        if (el.name or '').startswith('chrome') or el.kind == 'rect':
            continue
        if el.kind == 'text':
            role = _role(el)
            if role == 'skip':
                continue
            flat = _plain(el)
            if role in ('label', 'eyebrow') and LABEL_PREFIX.fullmatch(flat):
                pending = flat
                continue
            if pending:
                # A section divider is a number and then a title; merging the two would
                # cost the slide its heading. Only fold a label into ordinary text.
                if role in ('h2', 'h3'):
                    parts.append(f'<p class="ho-label">{esc(pending)}</p>')
                    pending = ''
                else:
                    parts.append(f'<p class="ho-choice"><b>{esc(pending)}</b> '
                                 f'{handout_runs(el.paras[0])}</p>')
                    pending = ''
                    if len(el.paras) > 1:
                        parts.append(handout_text(
                            Text(el.x, el.y, el.w, el.h, el.paras[1:], el.valign), role))
                    continue
            parts.append(handout_text(el, role))
        elif el.kind == 'image':
            src = copy_asset(el.src, assets_out)
            if src is not None:          # the still of a sketch not made yet: its link below stands in
                parts.append(f'<img src="{assets_rel}/{src}" alt="">')
        elif el.kind == 'figure':
            parts.append(f'<div class="ho-fig">{el.svg}</div>')
        elif el.kind == 'embed':
            parts.append(f'<p><a href="https://www.youtube.com/watch?v={esc(el.yt)}">'
                         f'Watch on YouTube &rsaquo;</a></p>')
        elif el.kind == 'sketch':
            hint = f' &mdash; {esc(el.hint)}' if el.hint else ''
            parts.append(f'<p><a href="sketches/{attr(el.name)}.html">Open the live sketch &rsaquo;</a>{hint}</p>')
        elif el.kind == 'exercise':
            # the widget itself is not duplicated — handout.js moves the one in the deck
            # into this slot, so ids and saved answers cannot diverge between the views
            parts.append(f'<div class="ho-ex" data-for="{attr(el.eid)}"></div>')
    if pending:
        parts.append(f'<p class="ho-label">{esc(pending)}</p>')
    if s.cp:
        label = s.cp['type'].replace('_', ' ')
        if s.report:
            parts.append(f'<p class="ho-cp"><a href="{attr(s.report)}" target="_blank" rel="noopener">ClassPoint &middot; {esc(label)} &mdash; see what the room answered</a></p>')
        else:
            parts.append(f'<p class="ho-cp">ClassPoint &middot; {esc(label)} &mdash; answer on the projector</p>')
    parts.append('</article>')
    return '\n'.join(parts)


def copy_asset(src, assets_out):
    """Copy an image into the html assets dir (downscaled to <= 1920px, jpeg where possible). None if it is missing."""
    from PIL import Image as PImage
    p = current().resolve_asset(src)
    if not p.exists():
        return None   # the still of a sketch not made yet; the callers show a labelled box instead
    assets_out.mkdir(parents=True, exist_ok=True)
    im = PImage.open(p)
    name = p.stem + ('.png' if im.mode in ('RGBA', 'LA', 'P') and p.suffix.lower() == '.png' and _has_alpha(im) else '.jpg')
    dst = assets_out / name
    if not dst.exists():
        im = im.convert('RGBA') if name.endswith('.png') else im.convert('RGB')
        if max(im.size) > 1920:
            im.thumbnail((1920, 1920))
        im.save(dst, quality=86, optimize=True) if name.endswith('.jpg') else im.save(dst, optimize=True)
    return name


def _has_alpha(im):
    return im.mode in ('RGBA', 'LA') and im.getextrema()[-1][0] < 255


PYODIDE_VERSION = '314.0.6'
PYODIDE_CDN = 'https://cdn.jsdelivr.net/pyodide/v{v}/full/'


def build_html(deck, out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    assets_out = out_dir / 'assets'
    slides = '\n'.join(html_slide(s, i + 1, out_dir, assets_out, 'assets') for i, s in enumerate(deck['slides']))
    # only ship the runtime if the deck can use it
    live = deck.get('console', True) and any(
        el.kind == 'exercise' for sl in deck['slides'] for el in sl.els + sl.html_only)
    handout = '\n'.join(handout_slide(sl, i + 1, assets_out, 'assets')
                        for i, sl in enumerate(deck['slides']))
    pyodide = ''
    if live:
        url = deck.get('pyodide_url') or PYODIDE_CDN.format(v=deck.get('pyodide_version', PYODIDE_VERSION))
        pyodide = PYODIDE_HTML.format(pyodide_url=url)
    css = CSS.replace('__CODE__', str(CODE)).replace('__CODE_SMALL__', str(CODE_SMALL))
    html = HTML_TMPL.format(title=esc(deck['title']), css=css, slides=slides,
                            handout=handout, pyodide=pyodide)
    (out_dir / 'index.html').write_text(html, encoding='utf-8')
    vendor = out_dir.parent / 'vendor'
    (vendor / 'fonts').mkdir(parents=True, exist_ok=True)
    for f in FONT_DIR.glob('*.ttf'):
        shutil.copy(f, vendor / 'fonts' / f.name)
    if live:
        for f in ('pyodide-console.js', 'pyodide-runtime.py'):
            shutil.copy(JS_DIR / f, vendor / f)
    shutil.copy(JS_DIR / 'handout.js', vendor / 'handout.js')   # every deck gets it
    return out_dir / 'index.html'


# ───────────────────────── PPTX backend ─────────────────────────
def build_pptx(deck, out_path: Path, footer: str):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    master = prs.slide_masters[0]
    blank = next(lo for lo in master.slide_layouts if lo.name == 'Blank')
    # keep a single layout: build.py installs the ait4x master + 8 layouts next to it
    for lo in list(master.slide_layouts):
        if lo is blank:
            continue
        rId = master.part.relate_to(lo.part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout')
        for el in list(master.element.sldLayoutIdLst):
            if el.rId == rId:
                master.element.sldLayoutIdLst.remove(el)
        master.part.drop_rel(rId)

    def rgb(c):
        return RGBColor.from_string(c.lstrip('#'))

    def add_text(slide, el):
        tb = slide.shapes.add_textbox(Emu(el.x * PX), Emu(el.y * PX), Emu(el.w * PX), Emu(el.h * PX))
        if el.name:
            tb.name = el.name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[el.valign]
        for i, p in enumerate(el.paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[p.align]
            size = max((r.size for r in p.runs), default=24)
            key = p.runs[0].font if p.runs else 'body'
            # css line-height is a multiple of font size; PowerPoint's is a multiple of the font's natural height
            para.line_spacing = round(p.lh * size / natural_lh(key, size), 3)
            if p.before:
                para.space_before = Pt(p.before * PT)
            if p.bullet:
                pPr = para._p.get_or_add_pPr()
                pPr.set('marL', str(48 * PX)); pPr.set('indent', str(-48 * PX))
                buClr = pPr.makeelement(qn('a:buClr'), {}); clr = buClr.makeelement(qn('a:srgbClr'), {'val': 'ED6D24'}); buClr.append(clr)
                buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Inter Black'})
                buChar = pPr.makeelement(qn('a:buChar'), {'char': '·'})
                for e in (buClr, buFont, buChar):
                    pPr.append(e)
            for r in p.runs:
                run = para.add_run()
                run.text = r.text.upper() if r.caps else r.text
                fam, bold, _w, _v, _c, _m = FONTS[r.font]
                f = run.font
                f.name = fam
                f.size = Pt(r.size * PT)
                f.bold = bold
                f.italic = r.italic or None
                f.color.rgb = rgb(r.color)
                rPr = run._r.get_or_add_rPr()
                if r.spc:
                    rPr.set('spc', str(int(round(r.spc * r.size * PT * 100))))
                if r.url:
                    run.hyperlink.address = r.url
        return tb

    def add_rect(slide, el):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(el.x * PX), Emu(el.y * PX), Emu(el.w * PX), Emu(el.h * PX))
        if el.name:
            sh.name = el.name
        if el.fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(el.fill)
        else:
            sh.fill.background()
        if el.stroke:
            sh.line.color.rgb = rgb(el.stroke); sh.line.width = Emu(el.stroke_w * PX)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def add_image(slide, path, el):
        from PIL import Image as PImage
        p = current().resolve_asset(path)
        if not p.exists():   # the still of a sketch not made yet: a labelled box where it will go
            add_rect(slide, Rect(el.x, el.y, el.w, el.h, PAPER))
            return add_text(slide, T(el.x, el.y, el.w, el.h, getattr(el, 'placeholder', '') or p.stem, 'monomed', 22, MUTED, lh=1.2, align='c', valign='m', spc=0.14, caps=True))
        im = PImage.open(p)
        iw, ih = im.size
        box_ar, img_ar = el.w / el.h, iw / ih
        x, y, w, h = el.x, el.y, el.w, el.h
        if getattr(el, 'fit', 'cover') == 'contain':
            if img_ar > box_ar:
                h = round(w / img_ar); y = el.y + (el.h - h) // 2
            else:
                w = round(h * img_ar); x = el.x + (el.w - w) // 2
            pic = slide.shapes.add_picture(str(p), Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
        else:
            pic = slide.shapes.add_picture(str(p), Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
            if img_ar > box_ar:      # too wide: crop left/right
                keep = box_ar / img_ar
                pic.crop_left = pic.crop_right = round((1 - keep) / 2, 4)
            elif img_ar < box_ar:    # too tall: crop top/bottom
                keep = img_ar / box_ar
                pic.crop_top = pic.crop_bottom = round((1 - keep) / 2, 4)
        if el.name:
            pic.name = el.name
        return pic

    for i, s in enumerate(deck['slides'], start=1):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = rgb(s.bg)
        for el in s.els:
            if el.kind == 'text':
                add_text(slide, el)
            elif el.kind == 'rect':
                add_rect(slide, el)
            elif el.kind == 'image':
                add_image(slide, el.src, el)
            elif el.kind == 'figure':
                add_image(slide, el.png, el)
            elif el.kind == 'exercise':
                for sub_el in exercise_static(el):
                    (add_rect if sub_el.kind == 'rect' else add_text)(slide, sub_el)
            elif el.kind == 'embed':
                if el.thumb:
                    add_image(slide, el.thumb, el)
                url = f'https://www.youtube.com/watch?v={el.yt}'
                add_text(slide, Text(el.x, el.y + el.h + 12, el.w, 36, [Para(runs(f'[{url}]({url})', 'mono', 22, MUTED), 'l', 1.3)]))
        notes = s.notes or ''
        slide.notes_slide.notes_text_frame.text = notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    _normalise_layout(out_path)
    # ClassPoint manifest for build.py (1-based slide numbers)
    manifest = {'_comment': f'Generated from the deck spec — slide number -> ClassPoint activity. Footer: {footer}'}
    for i, s in enumerate(deck['slides'], start=1):
        if s.cp:
            manifest[str(i)] = s.cp
    manifest_path = out_path.with_name(out_path.stem + '-activities.json')
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding='utf-8')
    return out_path, manifest_path


def _normalise_layout(pptx_path: Path):
    """python-pptx keeps the template's layout number (slideLayout7); build.py expects slideLayout1."""
    with zipfile.ZipFile(pptx_path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    layouts = [n for n in parts if re.match(r'ppt/slideLayouts/slideLayout\d+\.xml$', n)]
    assert len(layouts) == 1, layouts
    old = re.search(r'slideLayout(\d+)\.xml', layouts[0]).group(1)
    if old != '1':
        def ren(n):
            return re.sub(rf'slideLayout{old}\.xml', 'slideLayout1.xml', n)
        parts = {ren(n): d for n, d in parts.items()}
        for n in list(parts):
            if n.endswith('.rels') or n == '[Content_Types].xml':
                parts[n] = parts[n].decode('utf-8').replace(f'slideLayout{old}.xml', 'slideLayout1.xml').encode('utf-8')
    with zipfile.ZipFile(pptx_path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)


def run_classpoint_build(pptx_path: Path, manifest_path: Path, footer: str):
    from . import classpoint
    return classpoint.build(pptx_path, footer=footer, activities=manifest_path)


# ───────────────────────── PNG preview + overflow check ─────────────────────────
def build_png(deck, out_dir: Path, scale=0.5):
    from PIL import Image as PImage, ImageDraw
    out_dir.mkdir(parents=True, exist_ok=True)
    warnings = []
    files = []
    for i, s in enumerate(deck['slides'], start=1):
        im = PImage.new('RGB', (W, H), s.bg)
        d = ImageDraw.Draw(im)
        els = []
        for el in s.els:
            els.extend(exercise_static(el) if el.kind == 'exercise' else [el])
        for el in els:
            if el.kind == 'rect':
                if el.fill:
                    d.rectangle([el.x, el.y, el.x + el.w - 1, el.y + el.h - 1], fill=el.fill)
                if el.stroke:
                    d.rectangle([el.x, el.y, el.x + el.w - 1, el.y + el.h - 1], outline=el.stroke, width=el.stroke_w)
            elif el.kind in ('image', 'figure', 'embed'):
                src = el.src if el.kind == 'image' else (el.png if el.kind == 'figure' else el.thumb)
                if not src:
                    d.rectangle([el.x, el.y, el.x + el.w, el.y + el.h], fill='#111')
                    continue
                p = current().resolve_asset(src)
                if not p.exists():
                    warnings.append(f'slide {i:>2}: no still for the live sketch "{getattr(el, "placeholder", "") or p.stem}" — run `deckgen snap`')
                    d.rectangle([el.x, el.y, el.x + el.w, el.y + el.h], fill=PAPER)
                    d.text((el.x + 20, el.y + 20), (getattr(el, 'placeholder', '') or p.stem).upper(), font=pil_font('monomed', 22), fill=MUTED)
                    continue
                pim = PImage.open(p).convert('RGBA')
                fit = getattr(el, 'fit', 'cover')
                pim = _fit(pim, el.w, el.h, fit)
                ox = el.x + (el.w - pim.width) // 2; oy = el.y + (el.h - pim.height) // 2
                im.paste(pim, (ox, oy), pim)
            elif el.kind == 'text':
                th = text_height(el)
                if th > el.h + 2:
                    warnings.append(f'slide {i:>2}: text overflows box by {th - el.h:.0f}px — "{el.paras[0].runs[0].text[:50] if el.paras and el.paras[0].runs else ""}"')
                _draw_text(d, el, th)
        if s.cp:
            d.rectangle([1450, 908, 1800, 1000], outline=ORANGE, width=3)
            f = pil_font('monomed', 22)
            d.text((1470, 940), 'CLASSPOINT · ' + s.cp['type'].replace('_', ' ').upper(), font=f, fill=ORANGE)
        if scale != 1:
            im = im.resize((int(W * scale), int(H * scale)), PImage.LANCZOS)
        fp = out_dir / f'slide-{i:02d}.png'
        im.save(fp)
        files.append(fp)
    return files, warnings


def _fit(pim, w, h, fit):
    from PIL import Image as PImage
    iw, ih = pim.size
    if fit == 'contain':
        r = min(w / iw, h / ih)
        return pim.resize((max(1, round(iw * r)), max(1, round(ih * r))), PImage.LANCZOS)
    r = max(w / iw, h / ih)
    pim = pim.resize((max(1, round(iw * r)), max(1, round(ih * r))), PImage.LANCZOS)
    l = (pim.width - w) // 2; t = (pim.height - h) // 2
    return pim.crop((l, t, l + w, t + h))


def _draw_text(d, el, th):
    y = el.y
    if el.valign == 'm':
        y = el.y + (el.h - th) / 2
    elif el.valign == 'b':
        y = el.y + el.h - th
    for p in el.paras:
        size = max((r.size for r in p.runs), default=24)
        lines = wrap_para(p, el.w)
        y += p.before
        for li, line in enumerate(lines):
            lw = sum(run_width(r) for r, _ in line)
            x = el.x + {'l': 0, 'c': (el.w - lw) / 2, 'r': el.w - lw}[p.align]
            if p.bullet:
                x += 48
                if li == 0:
                    d.text((el.x, y + (p.lh * size - size) / 2), '·', font=pil_font('black', size), fill=ORANGE)
            base = y + (p.lh * size - natural_lh(p.runs[0].font if p.runs else 'body', size)) / 2
            for r, _ in line:
                f = pil_font(r.font, r.size)
                txt = r.text.upper() if r.caps else r.text
                if r.spc:
                    cx = x
                    for ch in txt:
                        d.text((cx, base), ch, font=f, fill=r.color)
                        cx += f.getlength(ch) + r.spc * r.size
                    x = cx
                else:
                    d.text((x, base), txt, font=f, fill=r.color)
                    x += f.getlength(txt)
            y += p.lh * size


def contact_sheet(files, out, cols=4, scale=0.5):
    from PIL import Image as PImage, ImageDraw
    ims = [PImage.open(f) for f in files]
    if not ims:
        return
    w, h = ims[0].size
    rows = (len(ims) + cols - 1) // cols
    sheet = PImage.new('RGB', (cols * (w + 12) + 12, rows * (h + 34) + 12), '#DDD')
    d = ImageDraw.Draw(sheet)
    for i, im in enumerate(ims):
        r, c = divmod(i, cols)
        x, y = 12 + c * (w + 12), 12 + r * (h + 34)
        sheet.paste(im, (x, y + 22))
        d.text((x, y + 4), f'{i + 1}', fill='#000', font=pil_font('monomed', 16))
    sheet.save(out, quality=80)
    return out


# ───────────────────────── live sketches ─────────────────────────
def write_sketch_page(el, out_dir: Path):
    """The standalone page of a live sketch: <deck>/sketches/<name>.html."""
    sk = out_dir / 'sketches'
    sk.mkdir(parents=True, exist_ok=True)
    sound = '<script src="../../vendor/p5/p5.sound.min.js"></script>' if el.sound else ''
    page = (SKETCH_TMPL.replace('@@NAME@@', esc(el.name)).replace('@@COURSE@@', esc(current().code))
            .replace('@@HINT@@', esc(el.hint or 'live p5.js')).replace('@@SOUND@@', sound)
            .replace('@@EXTRA@@', el.extra or '').replace('@@CODE@@', el.code)
            .replace('@@CW@@', str(el.cw)).replace('@@CH@@', str(el.ch)))
    path = sk / f'{el.name}.html'
    path.write_text(page, encoding='utf-8')
    return path


def sketches_of(deck):
    """Every live sketch of a deck, once each, in slide order."""
    seen = {}
    for s in deck['slides']:
        for el in s.els + s.html_only:
            if el.kind == 'sketch' and el.name not in seen:
                seen[el.name] = el
    return list(seen.values())


def node_env():
    """Environment for the node scripts in js/: NODE_PATH bridged to the course repo's node_modules (and
    the machine's global ones), since the scripts live in site-packages and would not see them."""
    env = dict(os.environ)
    roots = [p for p in (current().root / 'node_modules', Path('/opt/node22/lib/node_modules')) if p.is_dir()]
    if env.get('NODE_PATH'):
        roots.append(Path(env['NODE_PATH']))
    if roots:
        env['NODE_PATH'] = os.pathsep.join(str(p) for p in roots)
    return env


def playwright_env():
    """node_env() when node and the playwright package are here, else None."""
    if not shutil.which('node'):
        return None
    env = node_env()
    try:
        subprocess.run(['node', '-e', "require('playwright')"], check=True, env=env, capture_output=True, timeout=60)
    except Exception:
        return None
    return env


def snapshot_sketches(deck, name: str, force=False):
    """Screenshot every live sketch of a deck that needs a still into deck/assets/sketches/<sketch>.png
    (committed: the pptx and the PDF show it where the html deck runs the sketch). Skips the ones that
    exist unless force. Needs node + playwright + Chromium; without them the committed stills are used
    and missing ones are reported by the build. A page that makes no canvas (an error before
    createCanvas) writes no still: the committed one stays and a warning says so."""
    proj = current()
    todo = [el for el in sketches_of(deck) if el.twin and (force or not twin_png(el.name).exists())]
    if not todo:
        return []
    env = playwright_env()
    if env is None:
        return []
    # the pages load ../../vendor/p5/p5.min.js from _site/vendor: put it there, a clean checkout has not built the site
    shutil.copytree(SCAFFOLD_SITE / 'vendor' / 'p5', proj.site / 'vendor' / 'p5', dirs_exist_ok=True)
    out_dir = proj.site / name
    proj.sketches.mkdir(parents=True, exist_ok=True)
    args = []
    for el in todo:
        page = write_sketch_page(el, out_dir)
        args += [str(page), str(twin_png(el.name)), str(el.cw), str(el.ch)]
    started = time.time()
    subprocess.run(['node', str(JS_DIR / 'snap.js'), *args], env=env)   # exit 1 when a page made no canvas
    made = [twin_png(el.name) for el in todo if twin_png(el.name).exists() and twin_png(el.name).stat().st_mtime >= started]
    for el in todo:
        if twin_png(el.name) not in made:
            kept = 'the old still is kept' if twin_png(el.name).exists() else 'the build will show a placeholder'
            print(f'WARNING {name}: no still made for "{el.name}" (the page made no canvas: see the node output above); {kept}')
    return made


# ───────────────────────── orchestration ─────────────────────────
def build_pdf(index_html: Path, out_pdf: Path):
    """Print the html deck to a PDF (reveal.js print mode, Chromium via Playwright, js/pdf.js).
    No ClassPoint chips, video slides show their thumbnail. Needs node and the playwright
    package, plus a Chromium it can launch.

    pdf.js lives inside the installed package, so node resolves `require('playwright')`
    relative to site-packages and never sees the course repo's node_modules/. NODE_PATH is
    how that gets bridged — without it the PDF step fails on a machine where `npm install`
    put playwright exactly where the docs say to put it."""
    subprocess.run(['node', str(JS_DIR / 'pdf.js'), str(index_html), str(out_pdf)], check=True, env=node_env())
    return out_pdf


def build_all(deck, name: str, footer: str | None = None, do_pptx=True, do_html=True, do_png=True, do_pdf=True, snap=False):
    proj = current()
    SITE, EXPORT = proj.site, proj.export
    footer = proj.footer if footer is None else footer
    outputs = {}
    outputs['stills'] = snapshot_sketches(deck, name, force=snap)   # the twins of the live sketches, when node + playwright are here
    if do_html:
        outputs['html'] = build_html(deck, SITE / name)
        if do_pdf:  # the published, button-free version of the deck
            outputs['pdf'] = build_pdf(outputs['html'], SITE / name / deck.get('pdf', f'{name}.pdf'))
    if do_pptx:
        pptx_path, manifest = build_pptx(deck, EXPORT / f'{name}.pptx', footer)
        outputs['pptx'] = pptx_path
        outputs['manifest'] = manifest
        outputs['classpoint'] = run_classpoint_build(pptx_path, manifest, footer)
    if do_png:
        files, warnings = build_png(deck, EXPORT / 'preview' / name)
        outputs['png'] = files
        outputs['warnings'] = warnings
        for k in range(0, len(files), 16):
            contact_sheet(files[k:k + 16], EXPORT / 'preview' / f'{name}-sheet-{k // 16 + 1}.jpg')
    return outputs
